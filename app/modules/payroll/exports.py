"""Excel and PDF files of the salary report (and, later, the payments register).

The workbook builder mirrors ``web/app/people/_lib/xlsx.ts``: a title row, a muted meta line, a
tinted header row with an autofilter, number cells that stay numbers (``#,##0``), formula cells
with a cached value, bold total rows, several sheets per file.
"""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass
from io import BytesIO
from typing import Any, Literal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.core import bsdate
from app.core.export import generated_at_text, to_pdf
from app.modules.payroll import reports, rules
from app.modules.payroll.reports import ReportData, month_label

StyleName = Literal[
    "normal", "bold", "title", "muted", "header", "header_dark", "number", "number_bold", "total"
]


@dataclass(frozen=True)
class Styled:
    value: Any = None
    style: StyleName = "normal"
    #: Excel formula without the leading "=", with ``value`` as the cached result.
    formula: str | None = None


Cell = Any  # plain value or Styled


@dataclass
class Sheet:
    name: str
    rows: list[list[Cell]]
    widths: Sequence[float] = ()
    auto_filter: str | None = None


_HEADER_FILL = PatternFill("solid", fgColor="E7F1EF")
_DARK_FILL = PatternFill("solid", fgColor="0A5454")
_TOTAL_FILL = PatternFill("solid", fgColor="F1F3F0")
_HEADER_BORDER = Border(bottom=Side(style="thin", color="0A5454"))
_MUTED = "5B6670"
NUMBER_FORMAT = "#,##0"


def _sheet_title(name: str) -> str:
    cleaned = "".join("-" if ch in "[]:*?/\\" else ch for ch in name).strip()
    return (cleaned or "Sheet")[:31]


def _apply(cell, style: StyleName) -> None:
    if style == "title":
        cell.font = Font(bold=True, size=13)
    elif style == "bold":
        cell.font = Font(bold=True)
    elif style == "muted":
        cell.font = Font(color=_MUTED)
    elif style == "header":
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.border = _HEADER_BORDER
    elif style == "header_dark":
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _DARK_FILL
    elif style == "number":
        cell.number_format = NUMBER_FORMAT
        cell.alignment = Alignment(horizontal="right")
    elif style == "number_bold":
        cell.number_format = NUMBER_FORMAT
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="right")
    elif style == "total":
        cell.font = Font(bold=True)
        cell.fill = _TOTAL_FILL
        cell.number_format = NUMBER_FORMAT


def build_workbook(sheets: Sequence[Sheet]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in sheets:
        ws = wb.create_sheet(_sheet_title(sheet.name))
        for r, row in enumerate(sheet.rows, start=1):
            for c, raw in enumerate(row, start=1):
                styled = raw if isinstance(raw, Styled) else Styled(raw)
                if styled.formula:
                    cell = ws.cell(row=r, column=c, value=f"={styled.formula}")
                elif styled.value is None or styled.value == "":
                    cell = ws.cell(row=r, column=c)
                else:
                    cell = ws.cell(row=r, column=c, value=styled.value)
                    if isinstance(styled.value, str) and cell.data_type == "f":
                        cell.data_type = "s"
                if styled.style != "normal":
                    _apply(cell, styled.style)
        for i, width in enumerate(sheet.widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(width, 4)
        if sheet.auto_filter:
            ws.auto_filter.ref = sheet.auto_filter
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def header(labels: Sequence[str], dark: bool = False) -> list[Cell]:
    return [Styled(label, "header_dark" if dark else "header") for label in labels]


def num(value: int) -> Styled:
    return Styled(value, "number")


def total(value: int, formula: str | None = None) -> Styled:
    return Styled(value, "total", formula)


STATUS_LABEL = {"paid": "Paid", "part_paid": "Part paid", "pending": "Pending"}


def status_text(status: str, cleared_on: Any = None) -> str:
    """ "Cleared Kartik 8" once the month's due reached 0 (``cleared_on`` = AD date or a
    DateBothOut), else Part paid / Pending."""
    day = getattr(cleared_on, "ad", cleared_on)
    if day:
        return f"Cleared {bsdate.format_bs(bsdate.ad_to_bs(day), '%B %-d')}"
    return STATUS_LABEL[status]


# ------------------------------------------------------------------------ salary report


def salary_report_sheets(data: ReportData) -> list[Sheet]:
    out = data.out
    rows = data.ledgers
    all_branches = out.branch is None
    month = month_label(out.month.bs_year, out.month.bs_month)
    scope_line = f"{out.branch.name if out.branch else 'All branches'}" + (
        f" · {out.department.name}" if out.department else ""
    )

    HEADER = 4
    first = HEADER + 1
    last = HEADER + max(1, len(rows))
    columns = ["Emp ID", "Name", "Department"] + (["Branch"] if all_branches else []) + [
        "Pay type", "Paid days", "Days", "Gross", "Overtime", "Absence", "Advance paid",
        "Salary paid", "Net", "Paid", "Due", "Status",
    ]  # fmt: skip
    money_from = 3 + (1 if all_branches else 0) + 3  # 0-based index of "Gross"
    sheet_rows: list[list[Cell]] = [
        [Styled(f"TTK · {data.title}", "title")],
        [
            Styled(
                f"{len(rows)} staff · {scope_line} · net payroll Rs {out.totals.net:,}"
                f" · due Rs {out.totals.due:,}",
                "muted",
            )
        ],
        [],
        header(columns),
    ]
    for r in rows:
        e = r.employee
        line: list[Cell] = [e.code, e.full_name, e.department.name if e.department else ""]
        if all_branches:
            line.append(e.branch.name)
        line += [
            "Monthly" if r.structure.pay_type.value == "monthly" else "Daily wage",
            num(r.pay.paid_days),
            num(r.attendance.days_in_month),
            num(r.pay.earnings),
            num(r.pay.overtime_pay),
            num(r.pay.absence_deduction),
            num(r.settled.advance_paid),
            num(r.settled.salary_paid),
            Styled(r.pay.net, "number_bold"),
            num(r.settled.paid),
            Styled(r.settled.due, "number_bold"),
            status_text(r.settled.status, r.cleared_on),
        ]
        sheet_rows.append(line)
    if not rows:
        sheet_rows.append([Styled("Nobody on the sheet", "muted")])
    totals_line: list[Cell] = [Styled(f"Total · {len(rows)} staff", "total")] + [
        Styled(None, "total")
    ] * (money_from - 1)
    for offset, pick in enumerate(
        [
            lambda r: r.pay.earnings,
            lambda r: r.pay.overtime_pay,
            lambda r: r.pay.absence_deduction,
            lambda r: r.settled.advance_paid,
            lambda r: r.settled.salary_paid,
            lambda r: r.pay.net,
            lambda r: r.settled.paid,
            lambda r: r.settled.due,
        ]
    ):
        col = get_column_letter(money_from + offset + 1)
        totals_line.append(total(sum(pick(r) for r in rows), f"SUM({col}{first}:{col}{last})"))
    totals_line.append(Styled(None, "total"))
    sheet_rows.append(totals_line)
    widths = (
        [9, 24, 14]
        + ([14] if all_branches else [])
        + [11, 9, 7, 11, 10, 10, 12, 11, 11, 11, 11, 10]
    )
    payroll = Sheet(
        name="Payroll sheet",
        rows=sheet_rows,
        widths=widths,
        auto_filter=f"A{HEADER}:{get_column_letter(len(columns))}{last}",
    )

    dept_rows: list[list[Cell]] = [
        [Styled(f"By department · {month} · {scope_line}", "title")],
        header(["Department", "Staff", "Net payroll", "Paid", "Due"]),
    ]
    by_dept: dict[str, list[reports.ledgers.MonthLedger]] = {}
    for r in rows:
        by_dept.setdefault(
            r.employee.department.name if r.employee.department else "No department", []
        ).append(r)
    for name in sorted(by_dept):
        group = by_dept[name]
        dept_rows.append(
            [
                name,
                num(len(group)),
                num(sum(g.net for g in group)),
                num(sum(g.paid for g in group)),
                num(sum(g.due for g in group)),
            ]
        )
    d_first, d_last = 3, 2 + max(1, len(by_dept))
    dept_rows.append(
        [
            Styled("Total", "total"),
            total(len(rows), f"SUM(B{d_first}:B{d_last})"),
            total(out.totals.net, f"SUM(C{d_first}:C{d_last})"),
            total(out.totals.salary_paid + out.totals.advance_paid, f"SUM(D{d_first}:D{d_last})"),
            total(out.totals.due, f"SUM(E{d_first}:E{d_last})"),
        ]
    )
    by_department = Sheet(
        name="By department",
        rows=dept_rows,
        widths=[22, 8, 14, 14, 14],
        auto_filter=f"A2:E{d_last}",
    )

    sheets = [payroll]
    if all_branches:
        b_rows: list[list[Cell]] = [
            [
                Styled(
                    f"By branch · {month}"
                    + (f" · {out.department.name}" if out.department else ""),
                    "title",
                )
            ],
            header(["Branch", "Staff", "Net payroll", "Paid", "Due"]),
        ]
        for b in out.by_branch:
            b_rows.append([b.branch.name, num(b.staff), num(b.net), num(b.paid), num(b.due)])
        b_first, b_last = 3, 2 + max(1, len(out.by_branch))
        b_rows.append(
            [
                Styled("All branches", "total"),
                total(sum(b.staff for b in out.by_branch), f"SUM(B{b_first}:B{b_last})"),
                total(sum(b.net for b in out.by_branch), f"SUM(C{b_first}:C{b_last})"),
                total(sum(b.paid for b in out.by_branch), f"SUM(D{b_first}:D{b_last})"),
                total(sum(b.due for b in out.by_branch), f"SUM(E{b_first}:E{b_last})"),
            ]
        )
        sheets.append(
            Sheet(
                name="By branch",
                rows=b_rows,
                widths=[22, 8, 14, 14, 14],
                auto_filter=f"A2:E{b_last}",
            )
        )
    sheets.append(by_department)

    t = out.totals
    summary = Sheet(
        name="Summary",
        rows=[
            [Styled(f"Summary · {data.title}", "title")],
            ["Staff", num(t.staff)],
            [
                "Gross (incl. overtime)",
                num(t.gross),
                Styled(f"overtime Rs {t.overtime_pay:,} for {t.overtime_hours} h", "muted"),
            ],
            [
                "Absence deductions",
                num(t.absence_deduction),
                Styled(f"{t.absent_days} absent days", "muted"),
            ],
            [Styled("Net payable", "bold"), Styled(t.net, "number_bold")],
            [
                "Advances paid",
                num(t.advance_paid),
                Styled("part of the salary, already paid", "muted"),
            ],
            ["Salary paid", num(t.salary_paid)],
            [
                Styled("Still due", "bold"),
                Styled(t.due, "number_bold"),
                Styled(
                    f"{t.paid_count} paid · {t.part_paid_count} part paid"
                    f" · {t.pending_count} pending",
                    "muted",
                ),
            ],
            [],
            [
                Styled(
                    "Monthly: net = salary − salary ÷ 30 × absent days + overtime. "
                    "Daily: wage × days present + overtime. Saturdays are paid. No SSF.",
                    "muted",
                )
            ],
        ],
        widths=[26, 16, 60],
    )
    sheets.append(summary)
    return sheets


def salary_report_xlsx(data: ReportData) -> bytes:
    return build_workbook(salary_report_sheets(data))


def salary_report_pdf(data: ReportData, printed_by: str) -> bytes:
    out = data.out
    t = out.totals
    context = {
        "title": data.title,
        "subtitle": f"{len(data.ledgers)} staff · "
        f"{month_label(out.month.bs_year, out.month.bs_month)}",
        "printed_by": printed_by,
        "report": out,
        "rows": data.ledgers,
        "all_branches": out.branch is None,
        "status_label": STATUS_LABEL,
        "status_text": status_text,
        "legend_text": (
            "Monthly: net = salary − salary ÷ 30 × absent days + overtime · "
            "Daily: wage × days present + overtime · Saturdays paid · no SSF"
        ),
        "stats": [
            ("Gross", f"Rs {t.gross:,}", f"{t.staff} staff · overtime Rs {t.overtime_pay:,}"),
            ("Absence", f"− Rs {t.absence_deduction:,}", f"{t.absent_days} absent days"),
            ("Advances paid", f"Rs {t.advance_paid:,}", "part of salary"),
            (
                "Net payable",
                f"Rs {t.net:,}",
                f"{t.paid_count} paid · {t.part_paid_count} part · {t.pending_count} pending",
            ),
            ("Still due", f"Rs {t.due:,}", month_label(out.month.bs_year, out.month.bs_month)),
        ],
        "generated_at": generated_at_text(),
        "month_np": bsdate.format_bs(
            bsdate.BSDate(out.month.bs_year, out.month.bs_month, 1), "%B %Y", "ne"
        ),
    }
    return to_pdf("payroll/report.html", context, landscape=True)


# -------------------------------------------------------------------- payments register

from app.core.models import Employee  # noqa: E402
from app.modules.payroll.schemas import (  # noqa: E402
    MonthLedgerOut,
    PaymentOut,
    PaymentsRegisterOut,
    RegisterRowOut,
)


def _range_label(data: PaymentsRegisterOut) -> str:
    if data.from_bs and data.to_bs:
        return f"{data.from_bs} → {data.to_bs}"
    return "all dates"


def register_title(data: PaymentsRegisterOut) -> str:
    """``Payment report · {Branch|All branches} · {Department} · {range}``."""
    parts = [
        "Payment report",
        data.branch.name if data.branch else "All branches",
        data.department.name if data.department else None,
        _range_label(data),
    ]
    return " · ".join(p for p in parts if p)


def _slug(text: str) -> str:
    import re

    return re.sub(r"[^\w]+", "_", text).strip("_")


def register_file_name(data: PaymentsRegisterOut, ext: str) -> str:
    parts = [
        "Payment_report",
        _slug(data.branch.name) if data.branch else None,
        _slug(data.department.name) if data.department else None,
    ]
    span = f"{data.from_bs}_to_{data.to_bs}" if data.from_bs and data.to_bs else "all"
    return "_".join(p for p in [*parts, span] if p) + f".{ext}"


def _kind_label(p: PaymentOut) -> str:
    return "Salary" if p.kind == "salary" else "Advance"


def _split(rows: Sequence[RegisterRowOut]) -> tuple[int, int, int]:
    salary = sum(r.payment.amount for r in rows if r.payment.kind == "salary")
    advances = sum(r.payment.amount for r in rows if r.payment.kind == "advance")
    return salary, advances, salary + advances


def group_rows(rows: Sequence[RegisterRowOut], key) -> list[tuple[Any, list[RegisterRowOut]]]:
    """Consecutive groups in the order the rows come (the register sorts them)."""
    groups: list[tuple[Any, list[RegisterRowOut]]] = []
    for row in rows:
        k = key(row)
        if groups and groups[-1][0] == k:
            groups[-1][1].append(row)
        else:
            groups.append((k, [row]))
    return groups


def _dept_key(row: RegisterRowOut):
    return (
        (row.employee.department.id, row.employee.department.name)
        if row.employee.department
        else (None, "No department")
    )


def _branch_key(row: RegisterRowOut):
    """The branch that made the payment (not the employee's current branch)."""
    return (row.payment.branch.id, row.payment.branch.name)


def months_of(row: RegisterRowOut) -> str:
    """ "Ashoj 2083", or a salary payment group's split: "Bhadra 2083 3,000 · Ashoj 2083 5,000"."""
    if row.split and len(row.split) > 1:
        return " · ".join(
            f"{month_label(x.month.bs_year, x.month.bs_month)} {x.amount:,}" for x in row.split
        )
    return month_label(row.payment.for_year, row.payment.for_month)


def register_sheets(data: PaymentsRegisterOut) -> list[Sheet]:
    rows = data.rows
    all_branches = data.branch is None
    salary, advances, total_amount = _split(rows)
    rng = _range_label(data)
    month = month_label(data.totals.month.bs_year, data.totals.month.bs_month)

    HEADER = 4
    first = HEADER + 1
    last = HEADER + max(1, len(rows))
    # Date order reads oldest first on paper; the other sorts stay as the screen shows them.
    ordered = list(reversed(rows)) if data.sort == "date" else list(rows)
    columns = [
        "Date (BS)",
        "Emp ID",
        "Name",
        "Department",
        "Branch",
        "For month",
        "Type",
        "Paid by",
        "Reference / note",
        "Amount",
        "Recorded by",
        "Date (AD)",
    ]
    sheet_rows: list[list[Cell]] = [
        [Styled(f"TTK · {register_title(data)}", "title")],
        [
            Styled(
                f"{len(rows)} payments · salary Rs {salary:,} · advances Rs {advances:,}"
                f" · total Rs {total_amount:,}",
                "muted",
            )
        ],
        [],
        header(columns),
    ]
    for r in ordered:
        p, e = r.payment, r.employee
        sheet_rows.append(
            [
                p.date.bs,
                e.code,
                e.full_name,
                e.department.name if e.department else "",
                p.branch.name,
                months_of(r),
                _kind_label(p),
                p.payment_mode.name,
                " · ".join(x for x in (p.reference, p.note) if x),
                num(p.amount),
                p.recorded_by,
                p.date.ad,
            ]
        )
    if not rows:
        sheet_rows.append([Styled("No payments in this range", "muted")])
    sheet_rows.append(
        [Styled("Total", "total")]
        + [Styled(None, "total")] * 8
        + [
            total(total_amount, f"SUM(J{first}:J{last})"),
            Styled(f"=SUM(J{first}:J{last})", "muted"),
        ]
    )
    payments_sheet = Sheet(
        name="Payments",
        rows=sheet_rows,
        widths=[12, 9, 22, 14, 14, 14, 10, 16, 26, 12, 14, 12],
        auto_filter=f"A{HEADER}:L{last}",
    )

    def summary_sheet(
        name: str, label: str, groups: list[tuple[str, list[RegisterRowOut]]], title: str
    ) -> Sheet:
        g_first, g_last = 3, 2 + max(1, len(groups))
        body: list[list[Cell]] = [
            [Styled(title, "title")],
            header([label, "Salary", "Advances", "Total", "Count"]),
        ]
        for i, (gname, grows) in enumerate(groups):
            row = g_first + i
            s, a, t = _split(grows)
            body.append(
                [gname, num(s), num(a), Styled(t, "number_bold", f"B{row}+C{row}"), num(len(grows))]
            )
        if not groups:
            body.append([Styled("No payments in this range", "muted")])
        body.append(
            [
                Styled(f"Total · {len(groups)}", "total"),
                total(salary, f"SUM(B{g_first}:B{g_last})"),
                total(advances, f"SUM(C{g_first}:C{g_last})"),
                total(total_amount, f"SUM(D{g_first}:D{g_last})"),
                total(len(rows), f"SUM(E{g_first}:E{g_last})"),
            ]
        )
        return Sheet(name=name, rows=body, widths=[22, 14, 14, 14, 10], auto_filter=f"A2:E{g_last}")

    by_branch_groups = sorted(
        (
            (k[1], grows)
            for k, grows in group_rows(
                sorted(rows, key=lambda r: (_branch_key(r), r.payment.date.ad)), _branch_key
            )
        ),
        key=lambda g: g[0].lower(),
    )
    by_dept_groups = sorted(
        (
            (k[1], grows)
            for k, grows in group_rows(
                sorted(rows, key=lambda r: (_dept_key(r)[1], r.payment.date.ad)), _dept_key
            )
        ),
        key=lambda g: g[0].lower(),
    )
    sheets = [payments_sheet]
    if all_branches:
        sheets.append(
            summary_sheet(
                "By branch",
                "Branch",
                by_branch_groups,
                f"By branch · {rng}" + (f" · {data.department.name}" if data.department else ""),
            )
        )
    sheets.append(
        summary_sheet(
            "By department",
            "Department",
            by_dept_groups,
            f"By department · {rng}" + (f" · {data.department.name}" if data.department else ""),
        )
    )

    by_rows = data.by_employee
    b_first, b_last = 3, 2 + max(1, len(by_rows))
    by_employee_rows: list[list[Cell]] = [
        [Styled(f"By employee · {month}", "title")],
        header(
            [
                "Emp ID",
                "Name",
                "Department",
                "Branch",
                "Net pay",
                "Advances",
                "Salary paid",
                "Total paid",
                "Still due",
            ]
        ),
    ]
    for i, r in enumerate(by_rows):
        row = b_first + i
        by_employee_rows.append(
            [
                r.employee.code,
                r.employee.full_name,
                r.employee.department.name if r.employee.department else "",
                r.employee.branch.name,
                num(r.net),
                num(r.advance_paid),
                num(r.salary_paid),
                Styled(r.paid, "number_bold", f"F{row}+G{row}"),
                Styled(r.due, "number_bold", f"MAX(0,E{row}-H{row})"),
            ]
        )
    t = data.totals
    by_employee_rows.append(
        [
            Styled(f"Total · {len(by_rows)} staff", "total"),
            Styled(None, "total"),
            Styled(None, "total"),
            Styled(None, "total"),
            total(t.month_net, f"SUM(E{b_first}:E{b_last})"),
            total(t.advance_paid, f"SUM(F{b_first}:F{b_last})"),
            total(t.salary_paid, f"SUM(G{b_first}:G{b_last})"),
            total(t.salary_paid + t.advance_paid, f"SUM(H{b_first}:H{b_last})"),
            total(t.due, f"SUM(I{b_first}:I{b_last})"),
        ]  # fmt: skip
    )
    sheets.append(
        Sheet(
            name="By employee",
            rows=by_employee_rows,
            widths=[9, 22, 14, 14, 12, 12, 12, 12, 12],
            auto_filter=f"A2:I{b_last}",
        )
    )

    sheets.append(
        Sheet(
            name="Summary",
            rows=[
                [Styled(f"Summary · {month}", "title")],
                ["Net payroll", num(t.month_net), Styled(f"{t.staff} staff", "muted")],
                ["Salary paid", num(salary), Styled("in this range", "muted")],
                ["Advances paid", num(advances), Styled("in this range", "muted")],
                [Styled("Total paid", "bold"), Styled(total_amount, "number_bold", "B3+B4")],
                [
                    Styled(f"Still due · {month}", "bold"),
                    total(t.due),
                    Styled("net payroll − everything paid for the month", "muted"),
                ],
                [],
                [
                    Styled(
                        "Advances are part-payments of that month's salary; nothing is "
                        "recovered later. No SSF.",
                        "muted",
                    )
                ],
            ],
            widths=[26, 16, 44],
        )
    )

    f_first, f_last = 3, 2 + max(1, len(rows))
    filter_rows: list[list[Cell]] = [
        [
            Styled(
                f"All payments · sort & filter · {rng} · use the Name filter to pick employees",
                "title",
            )
        ],
        header(
            [
                "Date (BS)",
                "Emp ID",
                "Name",
                "Department",
                "Branch",
                "For month",
                "Type",
                "Paid by",
                "Amount",
                "Recorded by",
            ],
            dark=True,
        ),
    ]
    for r in rows:
        p, e = r.payment, r.employee
        filter_rows.append(
            [
                p.date.bs,
                e.code,
                e.full_name,
                e.department.name if e.department else "",
                p.branch.name,
                months_of(r),
                _kind_label(p),
                p.payment_mode.name,
                num(p.amount),
                p.recorded_by,
            ]
        )
    filter_rows.append(
        [Styled(f"Total (visible rows) · {len(rows)} rows", "total")]
        + [Styled(None, "total")] * 7
        + [
            total(total_amount, f"SUBTOTAL(109,I{f_first}:I{f_last})"),
            Styled(f"=SUBTOTAL(109,I{f_first}:I{f_last})", "muted"),
        ]
    )
    sheets.append(
        Sheet(
            name="Sort & filter",
            rows=filter_rows,
            widths=[12, 9, 22, 14, 14, 14, 10, 16, 12, 14],
            auto_filter=f"A2:J{f_last}",
        )
    )
    return sheets


def register_xlsx(data: PaymentsRegisterOut) -> bytes:
    return build_workbook(register_sheets(data))


def register_pdf(data: PaymentsRegisterOut, printed_by: str) -> bytes:
    rows = data.rows
    salary, advances, total_amount = _split(rows)
    t = data.totals
    month = month_label(t.month.bs_year, t.month.bs_month)
    all_branches = data.branch is None
    # Sort = branch: branch blocks. Sort = department: department blocks, nested inside branch
    # blocks when all branches are shown. Otherwise one flat block.
    blocks: list[dict[str, Any]] = []
    if data.sort == "branch":
        for (_, name), grows in group_rows(rows, _branch_key):
            s, a, tot = _split(grows)
            blocks.append(
                {
                    "kind": "branch",
                    "name": name,
                    "rows": grows,
                    "salary": s,
                    "advances": a,
                    "total": tot,
                    "departments": None,
                }
            )
    elif data.sort == "department" and all_branches:
        for (_, bname), brows in group_rows(rows, _branch_key):
            s, a, tot = _split(brows)
            depts = []
            for (_, dname), drows in group_rows(brows, _dept_key):
                ds, da, dt_ = _split(drows)
                depts.append(
                    {"name": dname, "rows": drows, "salary": ds, "advances": da, "total": dt_}
                )
            blocks.append(
                {
                    "kind": "branch",
                    "name": bname,
                    "rows": brows,
                    "salary": s,
                    "advances": a,
                    "total": tot,
                    "departments": depts,
                }
            )
    elif data.sort == "department":
        for (_, name), grows in group_rows(rows, _dept_key):
            s, a, tot = _split(grows)
            blocks.append(
                {
                    "kind": "department",
                    "name": name,
                    "rows": grows,
                    "salary": s,
                    "advances": a,
                    "total": tot,
                    "departments": None,
                }
            )
    else:
        blocks.append(
            {
                "kind": "flat",
                "name": "",
                "rows": list(rows),
                "salary": salary,
                "advances": advances,
                "total": total_amount,
                "departments": None,
            }
        )
    sort_label = {
        "date": "by date",
        "branch": "by branch",
        "department": "by department",
        "employee": "by employee",
        "amount": "by amount",
    }[data.sort]
    context = {
        "title": register_title(data),
        "subtitle": f"Salary & advance payments · {len(rows)} payments · {sort_label}",
        "printed_by": printed_by,
        "data": data,
        "blocks": blocks,
        "kind_label": _kind_label,
        "month_label": month_label,
        "stats": [
            ("Net payroll", f"Rs {t.month_net:,}", f"{t.staff} staff · {month}"),
            (
                "Salary paid",
                f"Rs {salary:,}",
                f"{sum(1 for r in rows if r.payment.kind == 'salary')} payments in range",
            ),
            (
                "Advances paid",
                f"Rs {advances:,}",
                f"{sum(1 for r in rows if r.payment.kind == 'advance')} advances in range",
            ),
            ("Still due", f"Rs {t.due:,}", f"for {month}"),
        ],
        "salary": salary,
        "advances": advances,
        "total_amount": total_amount,
        "month": month,
        "legend_text": "Advances are part-payments of that month's salary · amounts in rupees",
    }
    return to_pdf("payroll/register.html", context, landscape=True)


def receipt_pdf(
    employee: Employee,
    payment: PaymentOut,
    printed_by: str,
    group: Sequence[PaymentOut] | None = None,
) -> bytes:
    """``group`` = the rows of a salary payment group (oldest month first): one receipt for the
    total, split by month."""
    rows = list(group or [payment])
    split = (
        [(month_label(p.for_year, p.for_month), p.amount) for p in rows] if len(rows) > 1 else []
    )
    context = {
        "title": f"{'Salary payment' if payment.kind == 'salary' else 'Advance payment'}"
        f" · {payment.receipt_no}",
        "subtitle": f"{employee.full_name} · {employee.code}",
        "printed_by": printed_by,
        "employee": employee,
        "payment": payment,
        "amount": sum(p.amount for p in rows),
        "split": split,
        "month": month_label(payment.for_year, payment.for_month),
        "legend_text": (
            "An advance counts as part of that month's salary; it is not recovered later."
            if payment.kind == "advance"
            else "Signature: ______________________"
        ),
    }
    return to_pdf("payroll/receipt.html", context)


def payslip_pdf(employee: Employee, ledger: MonthLedgerOut, printed_by: str) -> bytes:
    per_day = round(ledger.rate / rules.PAY_DAYS)
    context = {
        "title": f"Payslip · {month_label(ledger.month.bs_year, ledger.month.bs_month)}"
        + (" (in progress)" if ledger.in_progress else ""),
        "subtitle": f"{employee.full_name} · {employee.code} · {employee.designation or ''}"
        f" · {employee.branch.name}",
        "printed_by": printed_by,
        "employee": employee,
        "ledger": ledger,
        "per_day": per_day,
        "status_label": STATUS_LABEL,
        "status_text": status_text,
        "month_label": month_label,
        "legend_text": (
            "Net = salary − salary ÷ 30 × absent days + overtime. Saturdays are paid. "
            "Advances are part of this month's salary."
            if ledger.pay_type == "monthly"
            else "Net = wage × (present + overtime days) + overtime hours × rate. "
            "Advances are part of this month's salary."
        ),
    }
    return to_pdf("payroll/payslip.html", context)
