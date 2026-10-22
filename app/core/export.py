"""Excel and PDF exports shared by every module.

* :func:`to_xlsx` ``(title, meta_lines, columns, rows) -> bytes`` - openpyxl; bold frozen header
  row, column widths from content, numbers stay numbers, text never becomes a formula.
* :func:`to_pdf` ``(template_name, context, landscape=False) -> bytes`` - WeasyPrint, A4, shared
  print stylesheet ``static/css/print.css`` with a running header (company, title, BS+AD date,
  printed by), footer (legend + "Page X of Y") and the embedded Noto Sans Devanagari font.

Print templates extend ``print/base.html``. ``context`` should carry ``title`` and
``printed_by``; ``company`` (default "TTK") and ``legend`` (list of ``(label, tone)``) are optional.
"""

from __future__ import annotations

from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.config import STATIC_DIR
from app.core import bsdate

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF_MEDIA_TYPE = "application/pdf"
PRINT_CSS = STATIC_DIR / "css" / "print.css"

_NUMBER_TYPES = (int, float, Decimal)


class PDFUnavailable(RuntimeError):
    """WeasyPrint could not be imported (on macOS: ``brew install pango``)."""


@dataclass(frozen=True)
class Column:
    key: str
    label: str
    width: float | None = None
    number_format: str | None = None  # e.g. "0", "#,##0.00", "yyyy-mm-dd"
    align: str | None = None  # "left" | "right" | "center"


ColumnSpec = str | tuple[str, str] | Column


def normalize_columns(columns: Sequence[ColumnSpec]) -> list[Column]:
    result: list[Column] = []
    for spec in columns:
        if isinstance(spec, Column):
            result.append(spec)
        elif isinstance(spec, str):
            result.append(Column(key=spec, label=spec))
        else:
            key, label = spec
            result.append(Column(key=key, label=label))
    return result


def _cell_value(row: Mapping[str, Any] | Sequence[Any], column: Column, index: int) -> Any:
    value = row.get(column.key) if isinstance(row, Mapping) else row[index]
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, datetime):
        # Excel has no timezones: store Kathmandu wall-clock time.
        if value.tzinfo is not None:
            value = value.astimezone(bsdate.KATHMANDU).replace(tzinfo=None)
        return value
    if isinstance(value, (*_NUMBER_TYPES, date)) or value is None:
        return value
    return str(value)


def _sheet_title(title: str) -> str:
    cleaned = "".join("-" if ch in "[]:*?/\\" else ch for ch in title).strip()
    return (cleaned or "Export")[:31]


def to_xlsx(
    title: str,
    meta_lines: Iterable[str],
    columns: Sequence[ColumnSpec],
    rows: Iterable[Mapping[str, Any] | Sequence[Any]],
) -> bytes:
    cols = normalize_columns(columns)
    wb = Workbook()
    ws = wb.active
    ws.title = _sheet_title(title)

    r = 1
    ws.cell(row=r, column=1, value=title).font = Font(bold=True, size=14)
    r += 1
    for line in meta_lines:
        ws.cell(row=r, column=1, value=line).font = Font(color="5B6670")
        r += 1
    r += 1  # blank line before the table

    header_row = r
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="E3F0EE")
    header_border = Border(bottom=Side(style="thin", color="0A5454"))
    for i, col in enumerate(cols, start=1):
        cell = ws.cell(row=r, column=i, value=col.label)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = header_border
        cell.alignment = Alignment(horizontal=col.align or "left", vertical="center")
    r += 1

    widths = [len(str(col.label)) for col in cols]
    for row in rows:
        for i, col in enumerate(cols, start=1):
            value = _cell_value(row, col, i - 1)
            cell = ws.cell(row=r, column=i, value=value)
            if isinstance(value, str) and cell.data_type == "f":
                cell.data_type = "s"  # text starting with "=" must stay text, never a formula
            if isinstance(value, _NUMBER_TYPES):
                cell.alignment = Alignment(horizontal=col.align or "right")
            elif col.align:
                cell.alignment = Alignment(horizontal=col.align)
            if col.number_format and isinstance(value, (*_NUMBER_TYPES, date, datetime)):
                cell.number_format = col.number_format
            widths[i - 1] = max(widths[i - 1], len("" if value is None else str(value)))
        r += 1

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    for i, col in enumerate(cols, start=1):
        width = col.width if col.width is not None else min(widths[i - 1] + 2, 60)
        ws.column_dimensions[get_column_letter(i)].width = max(width, 6)
    last_row = max(r - 1, header_row)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(cols))}{last_row}"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def to_pdf(template_name: str, context: Mapping[str, Any], landscape: bool = False) -> bytes:
    from app.core.templating import env  # late import: keeps export importable on its own

    full_context: dict[str, Any] = {
        "company": "TTK",
        "printed_by": "",
        "legend": [],
        "printed_at": bsdate.header_date(),
        "landscape": landscape,
        **context,
    }
    html = env.get_template(template_name).render(full_context)
    try:
        from weasyprint import CSS, HTML
        from weasyprint.text.fonts import FontConfiguration
    except (ImportError, OSError) as exc:  # not installed, or pango/cairo missing
        raise PDFUnavailable(
            "WeasyPrint cannot load its native libraries. On macOS run `brew install pango`."
        ) from exc
    # @font-face rules in CSS objects only apply when the same FontConfiguration is shared.
    font_config = FontConfiguration()
    stylesheets = [CSS(filename=str(PRINT_CSS), font_config=font_config)]
    if landscape:
        stylesheets.append(CSS(string="@page { size: A4 landscape; }", font_config=font_config))
    return HTML(string=html, base_url=f"{STATIC_DIR}/").write_pdf(
        stylesheets=stylesheets, font_config=font_config
    )
