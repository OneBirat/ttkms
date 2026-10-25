"""Voiding a payment: admin only, a reason, never twice, never undone. A voided payment keeps its
receipt number (never reused) and stays listed, but counts for nothing - balances, room, the
register's tiles and totals, exports and the payments-per-month contract all leave it out."""

from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app import contracts
from app.core import bsdate
from app.core.models import AuditLog, Branch, Department, Employee
from app.modules.admin.models import UserAccess
from app.modules.payroll.models import Payment, PayType, SalaryHistory

pytestmark = pytest.mark.anyio

TODAY = bsdate.today_bs()
CURRENT = (TODAY.year, TODAY.month)
PREVIOUS = bsdate.shift_month(*CURRENT, -1)


async def _anita(db):
    branch, dept = Branch(name="Main branch"), Department(name="Store")
    db.add_all([branch, dept])
    await db.flush()
    anita = Employee(
        full_name="Anita", branch_id=branch.id, department_id=dept.id, joined_ad=date(2026, 1, 1)
    )
    db.add(anita)
    await db.flush()
    db.add(
        SalaryHistory(
            employee_id=anita.id,
            pay_type=PayType.monthly,
            basic=30000,
            ot_rate=125,
            default_payment_mode_id=1,
            effective_ad=date(2026, 1, 1),
        )
    )
    await db.flush()
    return branch, anita


async def _csrf(client):
    return {"X-CSRF-Token": (await client.get("/api/auth/csrf")).json()["csrf_token"]}


async def _pay(client, h, anita, amount, month=PREVIOUS, kind="salary"):
    response = await client.post(
        "/api/payroll/payments",
        json={
            "employee_id": anita.id,
            "kind": kind,
            "amount": amount,
            "for_year": month[0],
            "for_month": month[1],
            "source": "account",
        },
        headers=h,
    )
    assert response.status_code == 201, response.text
    return response.json()["payment"]


async def _balance(client, anita, month=PREVIOUS):
    return (
        await client.get(
            "/api/payroll/balance",
            params={"employee_id": anita.id, "for_year": month[0], "for_month": month[1]},
        )
    ).json()


def _range():
    first = bsdate.BSDate(PREVIOUS[0], PREVIOUS[1], 1).isoformat()
    return {"from": first, "to": TODAY.isoformat()}


async def test_void_takes_a_payment_out_of_every_total(client, db, make_user, login):
    _, anita = await _anita(db)
    await make_user("admin", role="admin")
    await login("admin")
    h = await _csrf(client)
    kept = await _pay(client, h, anita, 5000)
    wrong = await _pay(client, h, anita, 7000)
    before = await _balance(client, anita)
    assert (before["paid"], before["due"]) == (12000, 18000)
    register = (await client.get("/api/payroll/payments", params=_range())).json()
    assert register["totals"]["count"] == 2 and register["totals"]["amount"] == 12000

    voided = await client.post(
        f"/api/payroll/payments/{wrong['id']}/void", json={"reason": "  Test   entry "}, headers=h
    )
    assert voided.status_code == 200, voided.text
    body = voided.json()
    assert body["voided"] is True and body["void_reason"] == "Test entry"
    assert body["voided_by"] == "Admin" and body["receipt_no"] == wrong["receipt_no"]

    after = await _balance(client, anita)
    assert (after["paid"], after["due"]) == (5000, 25000)
    register = (await client.get("/api/payroll/payments", params=_range())).json()
    assert register["totals"]["count"] == 1 and register["totals"]["amount"] == 5000
    assert [r["payment"]["id"] for r in register["rows"]] == [kept["id"]]
    assert register["show_voided"] is False
    shown = (
        await client.get("/api/payroll/payments", params={**_range(), "show_voided": "true"})
    ).json()
    assert sorted(r["payment"]["voided"] for r in shown["rows"]) == [False, True]
    assert shown["totals"]["amount"] == 5000  # listed, never counted
    assert sum(d["total"] for d in shown["by_department"]) == 5000

    # Exports skip it.
    xlsx = await client.get("/api/payroll/payments/export.xlsx", params=_range())
    cells = [
        str(c.value)
        for ws in load_workbook(BytesIO(xlsx.content)).worksheets
        for row in ws.iter_rows()
        for c in row
        if c.value is not None
    ]
    assert any(c.startswith("1 payments · salary Rs 5,000") for c in cells), cells[:3]

    # The payments-per-month contract (attendance badges) leaves it out.
    per_month = await contracts.get_payroll_payments().for_month(*PREVIOUS, [anita.id])
    assert [p.receipt_no for p in per_month[anita.id]] == [kept["receipt_no"]]

    # The ledger and transactions keep it, struck through, out of the sums.
    ledger = (
        await client.get(
            f"/api/payroll/employees/{anita.id}/ledger/{PREVIOUS[0]}-{PREVIOUS[1]:02d}"
        )
    ).json()
    assert ledger["paid"] == 5000 and ledger["salary_paid"] == 5000
    assert [p["voided"] for p in ledger["payments"]] == [False, True]
    tx = (await client.get(f"/api/payroll/employees/{anita.id}/transactions")).json()
    entry = next(e for e in tx["entries"] if e["id"] == f"pay-{wrong['id']}")
    assert entry["voided"] is True and entry["paid"] == 0
    assert entry["details"].startswith("Voided — Test entry")
    assert tx["totals"]["salary_payments"] == 5000 and tx["totals"]["salary_payment_count"] == 1

    # Receipt numbers are never reused.
    nxt = await _pay(client, h, anita, 100)
    assert int(nxt["receipt_no"].split("-")[1]) > int(wrong["receipt_no"].split("-")[1])

    audit = await db.scalar(select(AuditLog).where(AuditLog.action == "void"))
    assert audit.entity_id == str(wrong["id"]) and audit.data["after"]["reason"] == "Test entry"


async def test_void_frees_the_advance_room(client, db, make_user, login):
    _, anita = await _anita(db)
    await make_user("admin", role="admin")
    await login("admin")
    h = await _csrf(client)
    summary = (await client.get(f"/api/payroll/employees/{anita.id}/summary")).json()
    room = summary["room"]
    advance = await _pay(client, h, anita, room, CURRENT, "advance")
    assert (await client.get(f"/api/payroll/employees/{anita.id}/summary")).json()["room"] == 0
    await client.post(
        f"/api/payroll/payments/{advance['id']}/void", json={"reason": "Typo"}, headers=h
    )
    assert (await client.get(f"/api/payroll/employees/{anita.id}/summary")).json()["room"] == room


async def test_void_twice_and_reason_required(client, db, make_user, login):
    _, anita = await _anita(db)
    await make_user("admin", role="admin")
    await login("admin")
    h = await _csrf(client)
    p = await _pay(client, h, anita, 1000)
    empty = await client.post(
        f"/api/payroll/payments/{p['id']}/void", json={"reason": "  "}, headers=h
    )
    assert empty.status_code == 422
    missing = await client.post(f"/api/payroll/payments/{p['id']}/void", json={}, headers=h)
    assert missing.status_code == 422
    first = await client.post(
        f"/api/payroll/payments/{p['id']}/void", json={"reason": "Duplicate"}, headers=h
    )
    assert first.status_code == 200
    again = await client.post(
        f"/api/payroll/payments/{p['id']}/void", json={"reason": "Again"}, headers=h
    )
    assert again.status_code == 409 and again.json()["code"] == "already_voided"
    row = await db.get(Payment, p["id"])
    await db.refresh(row)
    assert row.void_reason == "Duplicate"
    unknown = await client.post(
        "/api/payroll/payments/999999/void", json={"reason": "x"}, headers=h
    )
    assert unknown.status_code == 404


async def test_only_admins_can_void(client, db, make_user, login):
    branch, anita = await _anita(db)
    await make_user("admin", role="admin")
    await login("admin")
    h = await _csrf(client)
    p = await _pay(client, h, anita, 1000)
    await client.post("/api/auth/logout", headers=h)
    # A manager with salary_advances switched on can pay, but not void.
    boss = await make_user("boss", role="manager")
    db.add(
        UserAccess(
            user_id=boss.id,
            branch_id=branch.id,
            access={
                "mark_attendance": True,
                "view_attendance_report": True,
                "manage_employees": True,
                "salary_advances": True,
            },
        )
    )
    await db.flush()
    await login("boss")
    h = await _csrf(client)
    denied = await client.post(
        f"/api/payroll/payments/{p['id']}/void", json={"reason": "No"}, headers=h
    )
    assert denied.status_code == 403
    row = await db.get(Payment, p["id"])
    assert row.voided_at is None
