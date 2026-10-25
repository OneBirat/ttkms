"""'Recorded by' (and 'Voided by', 'Printed by') name a person: the user's full name, else the
linked employee's name, else the username - never a raw phone number."""

from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.core import bsdate
from app.core.models import Branch, Department, Employee, User
from app.modules.payroll.models import PayType, SalaryHistory
from app.modules.payroll.services import user_label

pytestmark = pytest.mark.anyio

TODAY = bsdate.today_bs()
PREVIOUS = bsdate.shift_month(TODAY.year, TODAY.month, -1)


def test_user_label_prefers_names_and_never_shows_a_phone():
    anita = Employee(full_name="Anita Rai")
    assert user_label(User(id=1, username="ram", full_name="Ram Thapa")) == "Ram Thapa"
    # A phone typed as the name falls back to the linked employee, then the username.
    assert (
        user_label(User(id=2, username="9867301683", full_name="9867301683", employee=anita))
        == "Anita Rai"
    )
    assert user_label(User(id=3, username="sita", full_name=" ", phone="9800000000")) == "sita"
    # Only phone numbers to go on: a neutral label, not the number.
    phone_only = User(id=4, username="9867301683", full_name="+977 9867301683", phone="9867301683")
    assert user_label(phone_only) == "User 4"
    assert user_label(None) == ""


async def test_register_receipt_and_exports_show_the_full_name(client, db, make_user, login):
    branch, dept = Branch(name="Main branch"), Department(name="Store")
    db.add_all([branch, dept])
    await db.flush()
    anita = Employee(
        full_name="Anita", branch_id=branch.id, department_id=dept.id, joined_ad=date(2026, 1, 1)
    )
    db.add(anita)
    await db.flush()
    db.add(
        SalaryHistory(
            employee_id=anita.id,
            pay_type=PayType.monthly,
            basic=30000,
            ot_rate=125,
            default_payment_mode_id=1,
            effective_ad=date(2026, 1, 1),
        )
    )
    # The recorder logs in with a phone number as username.
    recorder = await make_user("9867301683", role="admin", full_name="Hari Prasad")
    recorder.phone = "9867301683"
    await db.flush()
    await login("9867301683")
    h = {"X-CSRF-Token": (await client.get("/api/auth/csrf")).json()["csrf_token"]}
    paid = await client.post(
        "/api/payroll/payments",
        json={
            "employee_id": anita.id,
            "kind": "salary",
            "amount": 5000,
            "for_year": PREVIOUS[0],
            "for_month": PREVIOUS[1],
            "source": "account",
        },
        headers=h,
    )
    assert paid.status_code == 201, paid.text
    assert paid.json()["payment"]["recorded_by"] == "Hari Prasad"

    params = {
        "from": bsdate.BSDate(PREVIOUS[0], PREVIOUS[1], 1).isoformat(),
        "to": TODAY.isoformat(),
        "branch_id": branch.id,
    }
    register = (await client.get("/api/payroll/payments", params=params)).json()
    assert [r["payment"]["recorded_by"] for r in register["rows"]] == ["Hari Prasad"]

    for path in ("ledger", "transactions"):
        body = (await client.get(f"/api/payroll/employees/{anita.id}/{path}")).json()
        assert "9867301683" not in str(body) and "Hari Prasad" in str(body), path

    xlsx = await client.get("/api/payroll/payments/export.xlsx", params=params)
    assert xlsx.status_code == 200
    values = [
        str(c)
        for ws in load_workbook(BytesIO(xlsx.content)).worksheets
        for row in ws.iter_rows(values_only=True)
        for c in row
        if c is not None
    ]
    assert "Hari Prasad" in values and not any("9867301683" in v for v in values)
