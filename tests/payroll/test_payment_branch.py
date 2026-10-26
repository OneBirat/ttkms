"""payments.branch_id: each payment keeps the branch the employee was in when it was made. After
a transfer, earlier payments stay in the old branch's Payment report, salary report totals and
exports; new payments go to the new branch."""

from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.models import AuditLog, Branch, Department, Employee
from app.modules.payroll.models import Payment, PayType, SalaryHistory

pytestmark = pytest.mark.anyio


async def _setup(db):
    old, new = Branch(name="Old branch"), Branch(name="New branch")
    store = Department(name="Store")
    db.add_all([old, new, store])
    await db.flush()
    anita = Employee(
        full_name="Anita", branch_id=old.id, department_id=store.id, joined_ad=date(2026, 1, 1)
    )
    db.add(anita)
    await db.flush()
    db.add(
        SalaryHistory(
            employee_id=anita.id,
            pay_type=PayType.monthly,
            basic=30000,
            ot_rate=125,
            default_payment_mode_id=1,
            effective_ad=date(2026, 1, 1),
        )
    )
    await db.flush()
    return old, new, anita


async def _admin(client, make_user, login):
    await make_user("admin", role="admin")
    await login("admin")
    return {"X-CSRF-Token": (await client.get("/api/auth/csrf")).json()["csrf_token"]}


async def _salary_for(client, h, anita, month, amount):
    """One salary payment for a completed BS month (works whatever the test "today" is)."""
    response = await client.post(
        "/api/payroll/payments",
        json={
            "employee_id": anita.id,
            "kind": "salary",
            "for_year": month["bs_year"],
            "for_month": month["bs_month"],
            "amount": amount,
        },
        headers=h,
    )
    assert response.status_code == 201, response.text
    return response.json()["payment"]


def _receipts(register) -> set[str]:
    return {r["payment"]["receipt_no"] for r in register["rows"]}


async def test_transfer_keeps_old_payments_in_the_old_branch(client, db, make_user, login):
    old, new, anita = await _setup(db)
    h = await _admin(client, make_user, login)

    # Before the transfer: 1,000 for the second month, then "Pay salary" 30,000 (clears the
    # oldest month).
    dues = (await client.get(f"/api/payroll/employees/{anita.id}/salary-dues")).json()
    oldest, second = dues["months"][0]["month"], dues["months"][1]["month"]
    before = await _salary_for(client, h, anita, second, 1000)
    salary = await client.post(
        "/api/payroll/salary-payments",
        json={"employee_id": anita.id, "amount": 30000, "payment_mode_id": 1},
        headers=h,
    )
    assert salary.status_code == 201, salary.text
    salary_receipt = salary.json()["receipt_no"]
    assert before["branch"]["id"] == old.id

    profile = (await client.get(f"/api/payroll/employees/{anita.id}")).json()
    moved = await client.patch(
        f"/api/payroll/employees/{anita.id}",
        json={
            "full_name": profile["full_name"],
            "phone": "9841000099",
            "branch_id": new.id,
            "department_id": profile["department"]["id"],
            "designation": "Store keeper",
            "joined_bs": profile["joined"]["bs"],
        },
        headers=h,
    )
    assert moved.status_code == 200, moved.text
    assert moved.json()["branch"]["id"] == new.id

    after = await _salary_for(client, h, anita, second, 500)
    assert after["branch"]["id"] == new.id
    stored = dict((await db.execute(select(Payment.receipt_no, Payment.branch_id))).all())
    assert stored == {
        before["receipt_no"]: old.id,
        salary_receipt: old.id,
        after["receipt_no"]: new.id,
    }
    audit = await db.scalar(
        select(AuditLog).where(AuditLog.action == "pay", AuditLog.entity_id == str(after["id"]))
    )
    assert audit.data["after"]["branch_id"] == new.id

    # Payment reports: by the branch that paid, not the employee's branch today.
    old_reg = (await client.get(f"/api/payroll/payments?branch_id={old.id}")).json()
    new_reg = (await client.get(f"/api/payroll/payments?branch_id={new.id}")).json()
    assert _receipts(old_reg) == {before["receipt_no"], salary_receipt}
    assert _receipts(new_reg) == {after["receipt_no"]}
    assert old_reg["totals"]["amount"] == 31000 and new_reg["totals"]["amount"] == 500
    every = (await client.get("/api/payroll/payments")).json()
    assert {b["branch"]["name"]: b["total"] for b in every["by_branch"]} == {
        "Old branch": 31000,
        "New branch": 500,
    }
    row = next(r for r in every["rows"] if r["payment"]["receipt_no"] == before["receipt_no"])
    assert row["payment"]["branch"]["name"] == "Old branch"
    assert row["employee"]["branch"]["name"] == "New branch"

    # Exports follow the same rule: the old branch's file keeps its rows, branch column = payer.
    xlsx = await client.get(f"/api/payroll/payments/export.xlsx?branch_id={old.id}")
    assert xlsx.status_code == 200, xlsx.text
    sheet = load_workbook(BytesIO(xlsx.content))["Payments"]
    values = [[c for c in r] for r in sheet.iter_rows(values_only=True)]
    header = next(i for i, r in enumerate(values) if r and r[0] == "Date (BS)")
    branch_col = values[header].index("Branch")
    data = [r for r in values[header + 1 :] if r and r[0] and str(r[0])[:2].isdigit()]
    assert len(data) == 2 and {r[branch_col] for r in data} == {"Old branch"}
    new_xlsx = await client.get(f"/api/payroll/payments/export.xlsx?branch_id={new.id}")
    new_sheet = load_workbook(BytesIO(new_xlsx.content))["Payments"]
    new_rows = [
        r for r in new_sheet.iter_rows(values_only=True) if r and r[0] and str(r[0])[:2].isdigit()
    ]
    assert len(new_rows) == 1

    # Salary report totals: the old branch still shows what it paid for the oldest month (and
    # Anita on that sheet); the new branch shows none of it.
    key = f"{oldest['bs_year']}-{oldest['bs_month']:02d}"
    old_report = (await client.get(f"/api/payroll/report/{key}?branch_id={old.id}")).json()
    new_report = (await client.get(f"/api/payroll/report/{key}?branch_id={new.id}")).json()
    assert [r["employee"]["full_name"] for r in old_report["rows"]] == ["Anita"]
    assert old_report["totals"]["salary_paid"] == 30000
    assert new_report["totals"]["salary_paid"] == 0
    month2 = f"{second['bs_year']}-{second['bs_month']:02d}"
    old_m2 = (await client.get(f"/api/payroll/report/{month2}?branch_id={old.id}")).json()
    new_m2 = (await client.get(f"/api/payroll/report/{month2}?branch_id={new.id}")).json()
    assert old_m2["totals"]["salary_paid"] == 1000
    assert new_m2["totals"]["salary_paid"] == 500
    everyone = (await client.get(f"/api/payroll/report/{month2}")).json()
    assert everyone["totals"]["salary_paid"] == 1500
    assert {b["branch"]["name"]: b["paid"] for b in everyone["by_branch"]} == {
        "Old branch": 1000,
        "New branch": 500,
    }
