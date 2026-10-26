"""Payments: the caps, for_month vs the payment date, mode checks, receipts, the register with
its filters, sorts and grouped totals, the account screens, exports and permissions."""

from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.core import bsdate
from app.core.models import AuditLog, Branch, Department, Employee
from app.modules.admin.models import PaymentMode, UserAccess
from app.modules.admin.permissions import effective_access
from app.modules.payroll.models import Payment, PayType, SalaryHistory

pytestmark = pytest.mark.anyio

TODAY = bsdate.today_bs()
CURRENT = (TODAY.year, TODAY.month)
PREVIOUS = bsdate.shift_month(*CURRENT, -1)
NEXT = bsdate.shift_month(*CURRENT, 1)


async def _csrf(client) -> str:
    return (await client.get("/api/auth/csrf")).json()["csrf_token"]


async def _setup(db):
    main, second = Branch(name="Main branch"), Branch(name="Branch 2")
    store, kitchen = Department(name="Store"), Department(name="Kitchen")
    db.add_all([main, second, store, kitchen])
    await db.flush()
    people = {}
    for name, branch, dept, pay, amount in [
        ("Anita", main, store, "monthly", 220000),
        ("Bikash", main, kitchen, "monthly", 320000),
        ("Ramesh", second, store, "daily", 650),
        ("Mina", second, kitchen, "monthly", 240000),
    ]:
        e = Employee(
            full_name=name, branch_id=branch.id, department_id=dept.id, joined_ad=date(2026, 1, 1)
        )
        db.add(e)
        await db.flush()
        db.add(
            SalaryHistory(
                employee_id=e.id,
                pay_type=PayType(pay),
                basic=amount,
                ot_rate=100,
                default_payment_mode_id=1,
                effective_ad=date(2026, 1, 1),
            )
        )
        people[name] = e
    bank = PaymentMode(
        type="bank",
        name="NIC Asia",
        details="A/c 4410",
        use_salary=True,
        use_advance=False,
        active=True,
    )
    off = PaymentMode(
        type="wallet", name="Old wallet", use_salary=True, use_advance=True, active=False
    )
    db.add_all([bank, off])
    await db.flush()
    return main, second, store, kitchen, people, bank, off


async def _admin(client, make_user, login):
    await make_user("admin", role="admin")
    await login("admin")
    return {"X-CSRF-Token": await _csrf(client)}


def _pay(employee, kind, amount, month=CURRENT, **extra):
    body = {
        "employee_id": employee.id,
        "kind": kind,
        "for_year": month[0],
        "for_month": month[1],
        "amount": amount,
        "source": "account",
    }
    body.update(extra)
    return body


async def _post(client, headers, body):
    return await client.post("/api/payroll/payments", json=body, headers=headers)


async def test_advance_is_capped_at_earned_so_far_minus_paid(client, db, make_user, login):
    *_, people, bank, off = await _setup(db)
    headers = await _admin(client, make_user, login)
    anita = people["Anita"]
    balance = (
        await client.get(
            f"/api/payroll/balance?employee_id={anita.id}&for_year={CURRENT[0]}&for_month={CURRENT[1]}"
        )
    ).json()
    room = balance["due"]
    assert (
        balance["earned"] > 0
        and room == balance["earned"] - balance["paid"]
        and balance["month_key"] == f"{CURRENT[0]}-{CURRENT[1]:02d}"
    )
    too_much = await _post(client, headers, _pay(anita, "advance", room + 1))
    assert (
        too_much.status_code == 422
        and too_much.json()["code"] == "over_room"
        and too_much.json()["cap"] == room
    )
    first = await _post(client, headers, _pay(anita, "advance", 1500))
    assert first.status_code == 201, first.text
    p = first.json()["payment"]
    assert (
        p["receipt_no"].startswith("ADV-")
        and p["kind"] == "advance"
        and p["payment_mode"]["name"] == "Cash"
    )
    assert (
        p["for_year"] == CURRENT[0]
        and p["for_month"] == CURRENT[1]
        and p["date"]["ad"] == bsdate.today_ad().isoformat()
    )
    assert first.json()["due_after"] == room - 1500
    # The room shrinks by what was paid; an advance never targets another month.
    second = await _post(client, headers, _pay(anita, "advance", room - 1500 + 1))
    assert second.status_code == 422 and second.json()["code"] == "over_room"
    exact = await _post(client, headers, _pay(anita, "advance", room - 1500))
    assert exact.status_code == 201 and exact.json()["due_after"] == 0
    past = await _post(client, headers, _pay(anita, "advance", 100, PREVIOUS))
    assert past.status_code == 422 and past.json()["field"] == "for_month"
    summary = (await client.get(f"/api/payroll/employees/{anita.id}/summary")).json()
    assert (
        summary["room"] == 0
        and summary["paid_this_month"] == room
        and summary["earned_so_far"] == balance["earned"]
    )
    assert [r["receipt_no"] for r in summary["recent"]][:1] == [
        exact.json()["payment"]["receipt_no"]
    ]


async def test_salary_caps_and_for_month_vs_payment_date(client, db, make_user, login, on_branch):
    *_, people, bank, off = await _setup(db)
    headers = await _admin(client, make_user, login)
    mina = people["Mina"]
    await on_branch("admin", mina.branch_id)  # Mina works in Branch 2
    prev = (
        await client.get(
            f"/api/payroll/balance?employee_id={mina.id}&for_year={PREVIOUS[0]}&for_month={PREVIOUS[1]}"
        )
    ).json()
    due = prev["due"]
    assert due == prev["earned"] and prev["status"] == "pending"
    over = await _post(client, headers, _pay(mina, "salary", due + 1, PREVIOUS))
    assert (
        over.status_code == 422 and over.json()["code"] == "over_due" and over.json()["cap"] == due
    )
    future = await _post(client, headers, _pay(mina, "salary", 100, NEXT))
    assert future.status_code == 422 and future.json()["field"] == "for_month"
    # Previous month's salary paid this month (today - on Kartik 1 that is the 1st) counts for
    # the previous month.
    paid_on = TODAY
    part = await _post(
        client,
        headers,
        _pay(
            mina,
            "salary",
            5000,
            PREVIOUS,
            date_bs=paid_on.isoformat(),
            payment_mode_id=bank.id,
            reference="CHQ 112",
        ),
    )
    assert part.status_code == 201, part.text
    p = part.json()["payment"]
    assert (p["for_year"], p["for_month"]) == PREVIOUS and p["date"]["bs"] == paid_on.isoformat()
    assert p["payment_mode"]["name"] == "NIC Asia" and p["reference"] == "CHQ 112"
    assert part.json()["due_after"] == due - 5000
    prev_after = (
        await client.get(
            f"/api/payroll/balance?employee_id={mina.id}&for_year={PREVIOUS[0]}&for_month={PREVIOUS[1]}"
        )
    ).json()
    assert (
        prev_after["paid"] == 5000
        and prev_after["due"] == due - 5000
        and prev_after["status"] == "part_paid"
    )
    this_month = (
        await client.get(
            f"/api/payroll/balance?employee_id={mina.id}&for_year={CURRENT[0]}&for_month={CURRENT[1]}"
        )
    ).json()
    assert this_month["paid"] == 0  # nothing counts for the running month
    rest = await _post(client, headers, _pay(mina, "salary", due - 5000, PREVIOUS))
    assert rest.status_code == 201 and rest.json()["due_after"] == 0
    assert (
        await client.get(
            f"/api/payroll/balance?employee_id={mina.id}&for_year={PREVIOUS[0]}&for_month={PREVIOUS[1]}"
        )
    ).json()["status"] == "paid"
    # The running month's salary is capped like an advance: earned so far − paid.
    running = (
        await client.get(
            f"/api/payroll/balance?employee_id={mina.id}&for_year={CURRENT[0]}&for_month={CURRENT[1]}"
        )
    ).json()
    adv = await _post(client, headers, _pay(mina, "advance", 1000))
    assert adv.status_code == 201
    too_much = await _post(client, headers, _pay(mina, "salary", running["due"] - 1000 + 1))
    assert too_much.status_code == 422 and too_much.json()["code"] == "over_room"
    ok = await _post(client, headers, _pay(mina, "salary", running["due"] - 1000))
    assert ok.status_code == 201 and ok.json()["due_after"] == 0
    # Receipts: one counter per kind, never repeating.
    receipts = (await db.execute(select(Payment.receipt_no))).scalars().all()
    assert len(receipts) == len(set(receipts))
    assert any(r.startswith("SAL-") for r in receipts) and any(
        r.startswith("ADV-") for r in receipts
    )
    audit = (await db.execute(select(AuditLog).where(AuditLog.entity == "payment"))).scalars().all()
    assert len(audit) == 4 and audit[0].data["after"]["receipt_no"] == p["receipt_no"]


async def test_mode_amount_and_date_validation(client, db, make_user, login):
    *_, people, bank, off = await _setup(db)
    headers = await _admin(client, make_user, login)
    anita = people["Anita"]
    for body, field in [
        (_pay(anita, "advance", 100, payment_mode_id=off.id), "payment_mode_id"),
        (
            _pay(anita, "advance", 100, payment_mode_id=bank.id),
            "payment_mode_id",
        ),  # bank: salary only
        (_pay(anita, "advance", 100, payment_mode_id=999), "payment_mode_id"),
        (_pay(anita, "advance", 0), "amount"),
        (_pay(anita, "advance", -5), "amount"),
        (_pay(anita, "advance", 100, date_bs="2099-01-01"), "date_bs"),
    ]:
        response = await _post(client, headers, body)
        assert response.status_code == 422, (body, response.text)
        assert response.json()["field"] == field
    assert (
        await _post(client, headers, _pay(anita, "advance", 100, for_month=13))
    ).status_code == 422
    assert (
        await _post(client, headers, _pay(anita, "salary", 100, payment_mode_id=bank.id))
    ).status_code == 201
    # A deactivated employee cannot be paid.
    assert (
        await client.post(f"/api/payroll/employees/{anita.id}/deactivate", headers=headers)
    ).status_code == 200
    assert (await _post(client, headers, _pay(anita, "advance", 100))).json()[
        "field"
    ] == "employee_id"


async def test_money_endpoints_need_salary_advances_and_stay_in_branch(
    client, db, make_user, login
):
    main, second, store, kitchen, people, bank, off = await _setup(db)
    anita, mina = people["Anita"], people["Mina"]
    boss = await make_user("boss", role="manager")
    db.add(UserAccess(user_id=boss.id, branch_id=main.id, access=None))
    await db.flush()
    await login("boss")
    headers = {"X-CSRF-Token": await _csrf(client)}
    assert (await client.get("/api/payroll/employees")).status_code == 200
    for path in (
        f"/api/payroll/employees/{anita.id}/ledger",
        f"/api/payroll/employees/{anita.id}/ledger/{CURRENT[0]}-{CURRENT[1]:02d}",
        f"/api/payroll/employees/{anita.id}/transactions",
        f"/api/payroll/employees/{anita.id}/summary",
        f"/api/payroll/balance?employee_id={anita.id}&for_year={CURRENT[0]}&for_month={CURRENT[1]}",
        "/api/payroll/payments",
        f"/api/payroll/employees/{anita.id}/payslips/{CURRENT[0]}-{CURRENT[1]:02d}.pdf",
    ):
        assert (await client.get(path)).status_code == 403, path
    assert (await _post(client, headers, _pay(anita, "advance", 100))).status_code == 403
    assert (
        await client.get(
            f"/api/payroll/employees/{anita.id}/attendance?bs_year={CURRENT[0]}&bs_month={CURRENT[1]}"
        )
    ).status_code == 200
    await client.post("/api/auth/logout", headers=headers)
    cashier = await make_user("cashier", role="manager")
    db.add(
        UserAccess(
            user_id=cashier.id,
            branch_id=second.id,
            access=effective_access("manager", {"salary_advances": True}),
        )
    )
    await db.flush()
    await login("cashier")
    headers = {"X-CSRF-Token": await _csrf(client)}
    assert (await _post(client, headers, _pay(mina, "advance", 100))).status_code == 201
    assert (await _post(client, headers, _pay(anita, "advance", 100))).status_code == 403
    assert (await client.get(f"/api/payroll/employees/{anita.id}/summary")).status_code == 403
    scoped = (await client.get(f"/api/payroll/payments?branch_id={main.id}")).json()
    assert scoped["branch"]["name"] == "Branch 2" and [
        r["employee"]["full_name"] for r in scoped["rows"]
    ] == ["Mina"]
    assert [b["branch"]["name"] for b in scoped["by_branch"]] == ["Branch 2"]
    await client.post("/api/auth/logout", headers=headers)
    await make_user("clerk", role="staff")
    await login("clerk")
    assert (await client.get("/api/payroll/payments")).status_code == 403


async def _seed_payments(client, headers, people, bank, on_branch):
    """Six payments across both branches and both departments, the previous month and this one.
    Each is made from the employee's branch (the admin switches to it; back to Main at the end).
    Ramesh is paid by the day, so he is marked present today first (nothing earned otherwise)."""
    first_of_month = bsdate.BSDate(CURRENT[0], CURRENT[1], 1).isoformat()
    await on_branch("admin", people["Ramesh"].branch_id)
    marked = await client.post(
        "/api/attendance/mark",
        json={"date_bs": TODAY.isoformat(), "employee_id": people["Ramesh"].id, "status": "P"},
        headers=headers,
    )
    assert marked.status_code == 200, marked.text
    made = []
    for name, kind, amount, month, extra in [
        ("Anita", "advance", 2000, CURRENT, {"date_bs": first_of_month}),
        ("Anita", "advance", 1000, CURRENT, {}),
        ("Bikash", "salary", 8000, PREVIOUS, {"payment_mode_id": bank.id, "reference": "TX1"}),
        ("Ramesh", "advance", 500, CURRENT, {}),
        ("Mina", "salary", 6000, PREVIOUS, {}),
        ("Mina", "advance", 1500, CURRENT, {}),
    ]:
        await on_branch("admin", people[name].branch_id)
        response = await _post(client, headers, _pay(people[name], kind, amount, month, **extra))
        assert response.status_code == 201, response.text
        made.append(response.json()["payment"])
    await on_branch("admin", people["Anita"].branch_id)
    return made


async def test_register_filters_sorts_and_grouped_totals(client, db, make_user, login, on_branch):
    main, second, store, kitchen, people, bank, off = await _setup(db)
    headers = await _admin(client, make_user, login)
    made = await _seed_payments(client, headers, people, bank, on_branch)
    first_of_month = bsdate.BSDate(CURRENT[0], CURRENT[1], 1).isoformat()
    rng = f"from={first_of_month}&to={TODAY.isoformat()}"
    # No branch asked for: the active branch (Main) - never every branch.
    active = (await client.get(f"/api/payroll/payments?{rng}")).json()
    assert active["branch"]["name"] == "Main branch" and active["sort"] == "date"
    assert len(active["rows"]) == 3
    dates = [r["payment"]["date"]["ad"] for r in active["rows"]]
    assert dates == sorted(dates, reverse=True)
    assert active["totals"]["count"] == 3 and active["totals"]["amount"] == 11000
    assert [(b["branch"]["name"], b["count"]) for b in active["by_branch"]] == [("Main branch", 3)]
    assert {d["department"]["name"]: d["count"] for d in active["by_department"]} == {
        "Kitchen": 1,
        "Store": 2,
    }
    assert sum(d["total"] for d in active["by_department"]) == 11000
    # Tiles: the settlement month is the previous month while anyone still owes for it.
    assert tuple(active["totals"]["month"].values()) == PREVIOUS
    assert active["totals"]["staff"] == 2
    assert active["totals"]["salary_paid"] == 8000 and active["totals"]["advance_paid"] == 0
    assert (
        len(active["by_employee"]) == 2
        and sum(e["due"] for e in active["by_employee"]) == active["totals"]["due"]
    )
    # Switching the active branch switches the register.
    await on_branch("admin", second.id)
    other = (await client.get(f"/api/payroll/payments?{rng}")).json()
    assert other["branch"]["name"] == "Branch 2" and other["totals"]["amount"] == 8000
    assert {r["employee"]["full_name"] for r in other["rows"]} == {"Ramesh", "Mina"}
    assert {d["department"]["name"]: d["count"] for d in other["by_department"]} == {
        "Kitchen": 2,
        "Store": 1,
    }
    # Filters inside the branch: type, department (chips keep the whole range), employee.
    advances = (await client.get(f"/api/payroll/payments?{rng}&type=advance")).json()
    assert [r["payment"]["kind"] for r in advances["rows"]] == ["advance", "advance"]
    kitchen_only = (
        await client.get(f"/api/payroll/payments?{rng}&department_id={kitchen.id}")
    ).json()
    assert [r["employee"]["full_name"] for r in kitchen_only["rows"]] == ["Mina", "Mina"]
    assert {d["department"]["name"]: d["count"] for d in kitchen_only["by_department"]} == {
        "Kitchen": 2,
        "Store": 1,
    }
    assert kitchen_only["totals"]["staff"] == 1 and kitchen_only["department"]["name"] == "Kitchen"
    await on_branch("admin", main.id)
    one = (await client.get(f"/api/payroll/payments?{rng}&employee_id={people['Anita'].id}")).json()
    assert len(one["rows"]) == 2
    # An explicit branch is still honoured for users who may see it.
    explicit = (await client.get(f"/api/payroll/payments?{rng}&branch_id={second.id}")).json()
    assert explicit["branch"]["name"] == "Branch 2" and len(explicit["rows"]) == 3
    # Sorts (one branch).
    by_branch = (await client.get(f"/api/payroll/payments?{rng}&sort=branch")).json()["rows"]
    assert [r["payment"]["branch"]["name"] for r in by_branch] == ["Main branch"] * 3
    flat = (await client.get(f"/api/payroll/payments?{rng}&sort=department")).json()["rows"]
    assert [r["employee"]["department"]["name"] for r in flat] == ["Kitchen", "Store", "Store"]
    by_employee = (await client.get(f"/api/payroll/payments?{rng}&sort=employee")).json()["rows"]
    names = [r["employee"]["full_name"] for r in by_employee]
    assert names == sorted(names)
    by_amount = (await client.get(f"/api/payroll/payments?{rng}&sort=amount")).json()["rows"]
    amounts = [r["payment"]["amount"] for r in by_amount]
    assert amounts == sorted(amounts, reverse=True)
    # Range: only the first of the month.
    narrow = (
        await client.get(f"/api/payroll/payments?from={first_of_month}&to={first_of_month}")
    ).json()
    on_the_first = [
        m["receipt_no"]
        for m in made
        if m["date"]["bs"] == first_of_month and m["branch"]["id"] == main.id
    ]
    assert sorted(r["payment"]["receipt_no"] for r in narrow["rows"]) == sorted(on_the_first)
    assert made[0]["receipt_no"] in on_the_first
    if TODAY.day > 1:
        assert on_the_first == [made[0]["receipt_no"]]
    last_month = bsdate.BSDate(PREVIOUS[0], PREVIOUS[1], 1).isoformat()
    bad = await client.get(f"/api/payroll/payments?from={first_of_month}&to={last_month}")
    assert bad.status_code == 422 and bad.json()["field"] == "from"
    await on_branch("admin", second.id)
    recent = (
        await client.get(f"/api/payroll/payments/recent?employee_id={people['Mina'].id}")
    ).json()
    assert [p["kind"] for p in recent] == ["advance", "salary"]


async def test_register_exports_and_receipt(client, db, make_user, login, on_branch):
    main, second, store, kitchen, people, bank, off = await _setup(db)
    headers = await _admin(client, make_user, login)
    made = await _seed_payments(client, headers, people, bank, on_branch)
    first_of_month = bsdate.BSDate(CURRENT[0], CURRENT[1], 1).isoformat()
    rng = f"from={first_of_month}&to={TODAY.isoformat()}"
    xlsx = await client.get(f"/api/payroll/payments/export.xlsx?{rng}&sort=department")
    assert xlsx.status_code == 200, xlsx.text
    # The active branch's file (Main): never "All branches", so no "By branch" sheet.
    assert xlsx.headers["content-disposition"].endswith(
        f'Payment_report_Main_branch_{first_of_month}_to_{TODAY.isoformat()}.xlsx"'
    )
    wb = load_workbook(BytesIO(xlsx.content))
    assert wb.sheetnames == ["Payments", "By department", "By employee", "Summary", "Sort & filter"]
    ws = wb["Payments"]
    assert ws["A1"].value.startswith("TTK · Payment report · Main branch · ")
    headers_row = [c.value for c in ws[4]]
    assert headers_row[3:5] == ["Department", "Branch"] and ws.auto_filter.ref == "A4:L7"
    assert ws.cell(row=8, column=10).value == "=SUM(J5:J7)"
    scoped = await client.get(
        f"/api/payroll/payments/export.xlsx?{rng}&branch_id={second.id}&department_id={kitchen.id}"
    )
    assert scoped.headers["content-disposition"].endswith(
        f'Payment_report_Branch_2_Kitchen_{first_of_month}_to_{TODAY.isoformat()}.xlsx"'
    )
    wb2 = load_workbook(BytesIO(scoped.content))
    assert "By branch" not in wb2.sheetnames
    assert wb2["Payments"]["A1"].value.startswith("TTK · Payment report · Branch 2 · Kitchen · ")
    pdf = await client.get(f"/api/payroll/payments/export.pdf?{rng}&sort=department")
    if pdf.status_code == 503:
        pytest.skip("WeasyPrint native libraries missing (brew install pango)")
    assert pdf.status_code == 200 and pdf.content[:5] == b"%PDF-"
    assert pdf.headers["content-disposition"].endswith('.pdf"')
    receipt = await client.get(f"/api/payroll/payments/{made[2]['id']}/receipt.pdf")
    assert receipt.status_code == 200 and receipt.content[:5] == b"%PDF-"
    assert receipt.headers["content-disposition"].endswith(f'receipt-{made[2]["receipt_no"]}.pdf"')
    await on_branch("admin", second.id)
    payslip = await client.get(
        f"/api/payroll/employees/{people['Mina'].id}/payslips/{PREVIOUS[0]}-{PREVIOUS[1]:02d}.pdf"
    )
    await on_branch("admin", main.id)
    assert payslip.status_code == 200 and payslip.content[:5] == b"%PDF-"
    # The register template groups branch → department with subtotals at both levels.
    from app.modules.payroll import exports, services
    from app.modules.payroll import payments as svc

    admin = await db.scalar(select(__import__("app.core.models", fromlist=["User"]).User))
    scope = await services.scope_for(db, admin, None)
    data = await svc.register(
        db,
        scope,
        from_bs=bsdate.BSDate.parse(first_of_month),
        to_bs=TODAY,
        kind=None,
        employee_id=None,
        department_id=None,
        sort="department",
    )
    from app.core.templating import env

    html = env.get_template("payroll/register.html").render(
        title=exports.register_title(data),
        data=data,
        blocks=[],
        kind_label=exports._kind_label,
        month_label=exports.month_label,
        stats=[],
        salary=0,
        advances=0,
        total_amount=0,
        month="",
        legend_text="",
        company="TTK",
        subtitle="",
        printed_by="",
        legend=[],
        printed_at=bsdate.header_date(),
        generated_at="",
        landscape=True,
    )
    assert "Payment report · Main branch · " in html
    assert exports.register_title(data).startswith("Payment report · Main branch · ")
    pdf_bytes = exports.register_pdf(data, "Test")
    assert pdf_bytes[:5] == b"%PDF-"


async def test_account_ledgers_and_transactions(client, db, make_user, login):
    *_, people, bank, off = await _setup(db)
    headers = await _admin(client, make_user, login)
    anita = people["Anita"]
    adv = await _post(client, headers, _pay(anita, "advance", 2000))
    assert adv.status_code == 201
    months = (await client.get(f"/api/payroll/employees/{anita.id}/ledger")).json()
    assert (
        len(months) == 6
        and tuple(months[0]["month"].values()) == CURRENT
        and months[0]["in_progress"] is True
    )
    assert months[0]["advance_paid"] == 2000 and months[0]["payments"][0]["kind"] == "advance"
    assert months[-1]["in_progress"] is False
    one = (
        await client.get(f"/api/payroll/employees/{anita.id}/ledger/{CURRENT[0]}-{CURRENT[1]:02d}")
    ).json()
    assert one["net"] == months[0]["net"] and one["status"] == "part_paid"
    assert (
        await client.get(f"/api/payroll/employees/{anita.id}/ledger/2070-01")
    ).status_code == 404
    tx = (await client.get(f"/api/payroll/employees/{anita.id}/transactions")).json()
    entries = tx["entries"]
    assert entries[0]["type"] == "salary" and entries[-1]["type"] == "advance"
    running = 0
    for e in entries:
        running += e["earned"] - e["paid"]
        assert e["balance"] == running
    assert (
        tx["totals"]["months"] == 6
        and tx["totals"]["advances"] == 2000
        and tx["totals"]["advance_count"] == 1
    )
    assert tx["totals"]["balance"] == sum(m["due"] for m in months)
    listed = (await client.get("/api/payroll/months")).json()
    assert len(listed) == 6 and tuple(listed[0].values()) == CURRENT
    attendance = (
        await client.get(
            f"/api/payroll/employees/{anita.id}/attendance?bs_year={CURRENT[0]}&bs_month={CURRENT[1]}"
        )
    ).json()
    assert attendance["summary"]["days_in_month"] == bsdate.days_in_month(*CURRENT)
    rows = (await client.get("/api/payroll/employees")).json()
    mine = next(r for r in rows if r["employee"]["id"] == anita.id)
    assert mine["balance_due"] == tx["totals"]["balance"]
