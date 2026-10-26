"""/api/payroll/report/{YYYY-MM}: scoping, department filter, by_branch, trend, Excel and PDF."""

from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.core import bsdate
from app.core.models import Branch, Department, Employee
from app.modules.admin.models import UserAccess
from app.modules.payroll.models import Payment, PaymentKind, PaymentSource, PayType, SalaryHistory

pytestmark = pytest.mark.anyio


async def _csrf(client) -> str:
    return (await client.get("/api/auth/csrf")).json()["csrf_token"]


async def _setup(db):
    """Two branches, two departments, four staff with salaries from January, one advance."""
    main, second = Branch(name="Main branch"), Branch(name="Branch 2")
    store, kitchen = Department(name="Store"), Department(name="Kitchen")
    db.add_all([main, second, store, kitchen])
    await db.flush()
    people = []
    for name, branch, dept, pay, amount in [
        ("Anita", main, store, "monthly", 220000),
        ("Bikash", main, kitchen, "monthly", 32000),
        ("Ramesh", second, store, "daily", 650),
        ("Mina", second, kitchen, "monthly", 24000),
    ]:
        e = Employee(
            full_name=name, branch_id=branch.id, department_id=dept.id, joined_ad=date(2026, 1, 1)
        )
        db.add(e)
        await db.flush()
        db.add(
            SalaryHistory(
                employee_id=e.id,
                pay_type=PayType(pay),
                basic=amount,
                ot_rate=100 if pay == "monthly" else 81,
                default_payment_mode_id=1,
                effective_ad=date(2026, 1, 1),
            )
        )
        people.append(e)
    today = bsdate.today_bs()
    db.add(
        Payment(
            receipt_no="ADV-0001",
            employee_id=people[0].id,
            branch_id=people[0].branch_id,
            kind=PaymentKind.advance,
            for_year=today.year,
            for_month=today.month,
            amount=2000,
            payment_mode_id=1,
            date=bsdate.today_ad(),
            source=PaymentSource.account,
        )
    )
    await db.flush()
    return main, second, store, kitchen, people


async def _admin(client, make_user, login):
    await make_user("admin", role="admin")
    await login("admin")


def _key(bs=None) -> str:
    bs = bs or bsdate.today_bs()
    return f"{bs.year}-{bs.month:02d}"


async def test_report_rows_totals_and_summaries(client, db, make_user, login, on_branch):
    main, second, store, kitchen, people = await _setup(db)
    await _admin(client, make_user, login)
    response = await client.get(f"/api/payroll/report/{_key()}")
    assert response.status_code == 200, response.text
    body = response.json()
    # No branch asked for: the admin's active branch (Main, the first created), never all.
    assert body["branch"]["name"] == "Main branch" and body["department"] is None
    assert [r["employee"]["full_name"] for r in body["rows"]] == ["Anita", "Bikash"]
    anita = body["rows"][0]["ledger"]
    assert (
        anita["pay_type"] == "monthly" and anita["rate"] == 220000 and anita["in_progress"] is True
    )
    assert (
        anita["advance_paid"] == 2000
        and anita["paid"] == 2000
        and anita["due"] == anita["net"] - 2000
    )
    assert anita["status"] == "part_paid" and anita["payments"][0]["kind"] == "advance"
    assert (
        anita["payments"][0]["month"] == body["month"]
        and anita["payments"][0]["payment_mode"]["name"] == "Cash"
    )
    assert (
        anita["net"] == anita["gross_salary"] - anita["absence_deduction"] + anita["overtime_pay"]
    )
    totals = body["totals"]
    assert totals["staff"] == 2 and totals["net"] == sum(r["ledger"]["net"] for r in body["rows"])
    assert (
        totals["due"] == sum(r["ledger"]["due"] for r in body["rows"])
        and totals["advance_paid"] == 2000
    )
    assert totals["paid_count"] + totals["part_paid_count"] + totals["pending_count"] == 2
    assert [b["branch"]["name"] for b in body["by_branch"]] == ["Main branch"]
    assert body["by_branch"][0]["staff"] == 2 and body["by_branch"][0]["net"] == totals["net"]
    assert body["by_branch"][0]["due"] == totals["due"]
    assert {d["department"]["name"]: d["staff"] for d in body["by_department"]} == {
        "Kitchen": 1,
        "Store": 1,
    }
    assert len(body["trend"]) == 5 and body["trend"][-1]["month"] == body["month"]
    assert body["all_staff"] == 2
    # Switching the active branch switches the report.
    await on_branch("admin", second.id)
    other = (await client.get(f"/api/payroll/report/{_key()}")).json()
    assert other["branch"]["name"] == "Branch 2"
    assert [r["employee"]["full_name"] for r in other["rows"]] == ["Mina", "Ramesh"]
    assert other["totals"]["advance_paid"] == 0


async def test_department_filter_and_branch_scoping(client, db, make_user, login):
    main, second, store, kitchen, people = await _setup(db)
    await _admin(client, make_user, login)
    kitchen_only = (
        await client.get(f"/api/payroll/report/{_key()}?department_id={kitchen.id}")
    ).json()
    assert [r["employee"]["full_name"] for r in kitchen_only["rows"]] == ["Bikash"]
    assert kitchen_only["all_staff"] == 2 and kitchen_only["department"]["name"] == "Kitchen"
    assert sum(b["staff"] for b in kitchen_only["by_branch"]) == 1
    branch_only = (await client.get(f"/api/payroll/report/{_key()}?branch_id={second.id}")).json()
    assert [r["employee"]["full_name"] for r in branch_only["rows"]] == ["Mina", "Ramesh"]
    assert branch_only["branch"]["name"] == "Branch 2" and branch_only["all_staff"] == 2
    assert [b["branch"]["name"] for b in branch_only["by_branch"]] == ["Branch 2"]
    assert {d["department"]["name"]: d["staff"] for d in branch_only["by_department"]} == {
        "Kitchen": 1,
        "Store": 1,
    }
    assert (await client.get(f"/api/payroll/report/{_key()}?branch_id=999")).status_code == 404
    assert (await client.get(f"/api/payroll/report/{_key()}?department_id=999")).status_code == 404
    assert (await client.get("/api/payroll/report/2099-01")).status_code == 404
    assert (await client.get("/api/payroll/report/2083-13")).status_code == 422
    # A manager without salary_advances gets 403; with the switch, only their own branch,
    # even with ?branch_id=.
    await client.post("/api/auth/logout", headers={"X-CSRF-Token": await _csrf(client)})
    boss = await make_user("boss", role="manager")
    db.add(UserAccess(user_id=boss.id, branch_id=second.id, access=None))
    await db.flush()
    await login("boss")
    assert (await client.get(f"/api/payroll/report/{_key()}")).status_code == 403
    await client.post("/api/auth/logout", headers={"X-CSRF-Token": await _csrf(client)})
    from app.modules.admin.permissions import effective_access

    cashier = await make_user("cashier", role="manager")
    db.add(
        UserAccess(
            user_id=cashier.id,
            branch_id=second.id,
            access=effective_access("manager", {"salary_advances": True}),
        )
    )
    await db.flush()
    await login("cashier")
    scoped = await client.get(f"/api/payroll/report/{_key()}?branch_id={main.id}")
    assert scoped.status_code == 200
    assert [r["employee"]["full_name"] for r in scoped.json()["rows"]] == ["Mina", "Ramesh"]
    assert scoped.json()["branch"]["name"] == "Branch 2"


async def test_excel_and_pdf_exports(client, db, make_user, login):
    main, second, store, kitchen, people = await _setup(db)
    await _admin(client, make_user, login)
    key = _key()
    xlsx = await client.get(f"/api/payroll/report/{key}.xlsx")
    assert xlsx.status_code == 200, xlsx.text
    # The active branch (Main): one branch, so no Branch column and no "By branch" sheet.
    assert xlsx.headers["content-disposition"].endswith(f'salary-{key}-main-branch.xlsx"')
    wb = load_workbook(BytesIO(xlsx.content))
    assert wb.sheetnames == ["Payroll sheet", "By department", "Summary"]
    ws = wb["Payroll sheet"]
    assert ws["A1"].value.startswith("TTK · Salary report · Main branch · ")
    headers = [c.value for c in ws[4]]
    assert headers[:4] == ["Emp ID", "Name", "Department", "Pay type"]
    assert ws.auto_filter.ref.startswith("A4:")
    net_col = headers.index("Net") + 1
    nets = [ws.cell(row=r, column=net_col).value for r in range(5, 7)]
    report = (await client.get(f"/api/payroll/report/{key}")).json()
    assert nets == [r["ledger"]["net"] for r in report["rows"]]
    total_row = 7
    assert ws.cell(row=total_row, column=net_col).value.startswith("=SUM(")
    # A branch + department: no "By branch" sheet, the names in the title and the file name.
    scoped = await client.get(
        f"/api/payroll/report/{key}.xlsx?branch_id={second.id}&department_id={kitchen.id}"
    )
    assert scoped.headers["content-disposition"].endswith(f'salary-{key}-branch-2-kitchen.xlsx"')
    wb2 = load_workbook(BytesIO(scoped.content))
    assert wb2.sheetnames == ["Payroll sheet", "By department", "Summary"]
    assert wb2["Payroll sheet"]["A1"].value.startswith(
        "TTK · Salary report · Branch 2 · Kitchen · "
    )
    # TDS stays in the data (0) but no export shows it.
    cells = [
        str(c)
        for book in (wb, wb2)
        for sheet in book.worksheets
        for row in sheet.iter_rows(values_only=True)
        for c in row
        if c is not None
    ]
    assert not any("TDS" in c for c in cells)
    assert [c.value for c in wb2["Payroll sheet"][4]][:4] == [
        "Emp ID",
        "Name",
        "Department",
        "Pay type",
    ]
    pdf = await client.get(f"/api/payroll/report/{key}.pdf?department_id={store.id}")
    if pdf.status_code == 503:
        pytest.skip("WeasyPrint native libraries missing (brew install pango)")
    assert pdf.status_code == 200 and pdf.content[:5] == b"%PDF-"
    assert pdf.headers["content-disposition"].endswith(f'salary-{key}-main-branch-store.pdf"')
    # The rendered HTML carries the title rule.
    from app.core.templating import env
    from app.modules.payroll import exports, reports, services

    scope = await services.scope_for(
        db,
        (
            await db.scalar(
                __import__("sqlalchemy").select(
                    __import__("app.core.models", fromlist=["User"]).User
                )
            )
        ),
        None,
    )
    data = await reports.salary_report(
        db, scope, bsdate.today_bs().year, bsdate.today_bs().month, store.id
    )
    assert data.title.startswith("Salary report · Main branch · Store · ")
    html = env.get_template("payroll/report.html").render(
        title=data.title,
        report=data.out,
        rows=data.ledgers,
        all_branches=False,
        status_label=exports.STATUS_LABEL,
        status_text=exports.status_text,
        stats=[],
        month_np="",
        company="TTK",
        subtitle="",
        printed_by="",
        legend=[],
        legend_text="",
        printed_at=bsdate.header_date(),
        generated_at="",
        landscape=True,
    )
    assert "Salary report · Main branch · Store" in html and "Anita" in html
    assert "Ramesh" not in html  # Branch 2


def test_pdf_templates_never_show_tds():
    from pathlib import Path

    folder = Path(__file__).resolve().parents[2] / "app/modules/payroll/templates/payroll"
    for template in folder.glob("*.html"):
        assert "TDS" not in template.read_text() and ".tds" not in template.read_text(), template
