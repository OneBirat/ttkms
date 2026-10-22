import zipfile
from datetime import UTC, date, datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.core.export import (
    PDF_MEDIA_TYPE,
    XLSX_MEDIA_TYPE,
    Column,
    PDFUnavailable,
    to_pdf,
    to_xlsx,
)

pytestmark = pytest.mark.anyio

COLUMNS = [
    Column("name", "Item"),
    Column("name_np", "नाम"),
    Column("qty", "Qty", number_format="0"),
]
ROWS = [
    {"name": "Pressure cooker", "name_np": "प्रेसर कुकर", "qty": 24},
    {"name": "Gas stove", "name_np": "ग्यास चुलो", "qty": 3},
]


def _weasyprint_available() -> bool:
    try:
        import weasyprint  # noqa: F401
    except (ImportError, OSError):
        return False
    return True


def test_to_xlsx_layout():
    data = to_xlsx("Stock items", ["TTK", "Printed by Admin"], COLUMNS, ROWS)
    wb = load_workbook(BytesIO(data))
    ws = wb.active
    assert ws.title == "Stock items"
    assert ws["A1"].value == "Stock items" and ws["A1"].font.bold
    assert ws["A2"].value == "TTK" and ws["A3"].value == "Printed by Admin"
    header_row = 5
    assert [c.value for c in ws[header_row]] == ["Item", "नाम", "Qty"]
    assert all(c.font.bold for c in ws[header_row])
    assert ws.freeze_panes == "A6"
    assert ws.cell(row=6, column=3).value == 24
    assert ws.cell(row=6, column=3).data_type == "n", "numbers stay numbers"
    assert ws.cell(row=6, column=3).number_format == "0"
    assert ws.cell(row=7, column=2).value == "ग्यास चुलो"
    assert ws.column_dimensions["A"].width >= len("Pressure cooker")
    assert ws.auto_filter.ref == "A5:C7"


def test_to_xlsx_accepts_tuples_and_sequences():
    data = to_xlsx("T", [], [("a", "A"), "B"], [(1, "x"), (2, None)])
    ws = load_workbook(BytesIO(data)).active
    assert ws["A3"].value == "A" and ws["B3"].value == "B"  # no meta lines: header on row 3
    assert ws["A4"].value == 1 and ws["B5"].value is None
    assert ws.title == "T"


def test_sheet_title_is_sanitised():
    data = to_xlsx(
        "Report: 2083/07 [draft] with a very long title that exceeds limits", [], ["A"], []
    )
    ws = load_workbook(BytesIO(data)).active
    assert len(ws.title) <= 31 and not any(ch in ws.title for ch in "[]:*?/\\")


def test_text_starting_with_equals_is_not_a_formula():
    data = to_xlsx("T", [], ["note"], [{"note": '=HYPERLINK("http://evil","x")'}])
    with zipfile.ZipFile(BytesIO(data)) as zf:
        sheet_xml = zf.read("xl/worksheets/sheet1.xml").decode()
    assert "<f>" not in sheet_xml, "must be stored as text, never as a live formula"
    ws = load_workbook(BytesIO(data)).active
    assert ws["A4"].value == '=HYPERLINK("http://evil","x")'


def test_timezone_aware_datetimes_and_dates_are_written():
    when = datetime(2026, 10, 20, 3, 30, tzinfo=UTC)  # 09:15 in Kathmandu
    data = to_xlsx(
        "T",
        [],
        [
            Column("at", "At", number_format="yyyy-mm-dd hh:mm"),
            Column("on", "On", number_format="yyyy-mm-dd"),
        ],
        [{"at": when, "on": date(2026, 10, 20)}],
    )  # fmt: skip
    ws = load_workbook(BytesIO(data)).active
    assert ws["A4"].value == datetime(2026, 10, 20, 9, 15)
    assert ws["A4"].number_format == "yyyy-mm-dd hh:mm"
    assert ws["B4"].value == datetime(2026, 10, 20) and ws["B4"].number_format == "yyyy-mm-dd"


def test_to_pdf_renders_a4_with_devanagari():
    if not _weasyprint_available():
        pytest.skip("WeasyPrint native libraries missing (brew install pango)")
    pdf = to_pdf(
        "print/export_demo.html",
        {"title": "Demo", "printed_by": "Admin", "columns": COLUMNS, "rows": ROWS},
    )
    assert pdf.startswith(b"%PDF")
    assert len(pdf) > 2000
    landscape = to_pdf(
        "print/export_demo.html",
        {"title": "Demo", "columns": COLUMNS, "rows": ROWS},
        landscape=True,
    )
    assert landscape.startswith(b"%PDF")


def test_to_pdf_without_native_libraries_raises_clear_error():
    if _weasyprint_available():
        pytest.skip("WeasyPrint is importable here; the fallback path is covered elsewhere")
    with pytest.raises(PDFUnavailable, match="pango"):
        to_pdf("print/export_demo.html", {"title": "x", "columns": [], "rows": []})


async def test_export_demo_routes(client, make_user, login):
    await make_user("admin", full_name="Site Admin")
    assert (await client.get("/api/core/export-demo.xlsx")).status_code == 401
    await login("admin")

    xlsx = await client.get("/api/core/export-demo.xlsx")
    assert xlsx.status_code == 200
    assert xlsx.headers["content-type"].startswith(XLSX_MEDIA_TYPE)
    assert "attachment" in xlsx.headers["content-disposition"]
    ws = load_workbook(BytesIO(xlsx.content)).active
    assert ws["A1"].value.startswith("Export demo")
    assert ws["A4"].value == "Printed by Site Admin"

    pdf = await client.get("/api/core/export-demo.pdf")
    if pdf.status_code == 503:
        pytest.skip(f"PDF export unavailable: {pdf.json()['detail']}")
    assert pdf.status_code == 200
    assert pdf.headers["content-type"].startswith(PDF_MEDIA_TYPE)
    assert pdf.content.startswith(b"%PDF")
